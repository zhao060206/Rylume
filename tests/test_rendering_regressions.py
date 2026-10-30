"""Regression tests for SVG/PPT rendering and audit edge cases."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import app_server
import compose
import framework
import model_review
import layout_contract
import pipeline
import ppt_native
import review
import archetypes_render
import assist_guide


class SafetyTests(unittest.TestCase):
    def test_event_handler_is_rejected(self):
        svg = (b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 200 100">'
               b'<rect x="1" y="1" width="50" height="20" onclick="alert(1)"/></svg>')
        with self.assertRaises(RuntimeError):
            pipeline._check_svg(svg)

    def test_model_colors_cannot_break_attributes(self):
        bad = '#fff" onclick="alert(1)'
        scene = {
            "canvas": [400, 200], "vector_glyphs": True,
            "palette": {"primary": bad, "text": bad},
            "entities": [{"id": "a", "label": "A", "bbox": [40, 40, 120, 90]}],
            "connectors": [], "labels": [],
        }
        glyph = framework._vector_glyph("A", 50, 50, 78, bad)
        asset = ('<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100">'
                 + glyph + '</svg>').encode()
        svg = compose.compose_svg(scene, {"a": (asset, 100.0, 100.0)}, "Times New Roman")
        self.assertNotIn(b"onclick", svg)
        self.assertNotIn(b'alert(1)"', svg)


class SignalLayoutTests(unittest.TestCase):
    def test_spatial_hints_are_parsed(self):
        text = ("空间布局：采用左到右的轻微 S 形流程。左侧 0–25% 为 Extracellular space；"
                "25–40% 为 Plasma membrane；40–70% 为 Cytoplasm；70–100% 为 Nucleus。")
        layout, regions = archetypes_render._layout_hints(text)
        self.assertEqual(layout, "horizontal_s")
        self.assertEqual([r["label"] for r in regions],
                         ["Extracellular space", "Plasma membrane", "Cytoplasm", "Nucleus"])

    def test_horizontal_s_keeps_all_five_stages(self):
        spec = {
            "title": "EV Signaling", "subtitle": "Left to right",
            "layout": "horizontal_s",
            "regions": [{"start": 0, "end": 25, "label": "Extracellular space"},
                        {"start": 25, "end": 40, "label": "Plasma membrane"}],
            "stages": [
                {"label": "EV secretion", "caption": "EVs release ligands.", "color": "#7A5C3E"},
                {"label": "Receptor binding", "caption": "Ligands bind the receptor.", "color": "#337AB7"},
                {"label": "Kinase cascade", "caption": "Receptor activates kinases.", "color": "#2E8B7A"},
                {"label": "Nuclear entry", "caption": "Factors enter the nucleus.", "color": "#B7791F"},
                {"label": "Matrix synthesis", "caption": "Chondrocytes restore matrix.", "color": "#6B46C1"},
            ],
        }
        svg = archetypes_render.RENDERERS["signal_pathway"](
            spec, {}, [1440, 1080], "rect", archetypes_render._palette({}, "signal_pathway"))
        pipeline._check_svg(svg, allow_image=True)
        for stage in spec["stages"]:
            self.assertIn(stage["label"].encode(), svg)
        self.assertEqual(review.findings(svg, None), [])

    def test_vertical_request_keeps_vertical_layout(self):
        layout, _ = archetypes_render._layout_hints("自上而下单向通路，纵向排列。")
        self.assertEqual(layout, "vertical")


class PlateSceneTests(unittest.TestCase):
    SPEC = {
        "title": "EV Signaling", "subtitle": "Left to right",
        "stages": [
            {"label": "Ligand release", "caption": "Ligands enter the space.", "color": "#7A5C3E"},
            {"label": "Receptor binding", "caption": "Ligands bind the receptor.", "color": "#337AB7"},
            {"label": "Kinase cascade", "caption": "Kinases relay signaling.", "color": "#2E8B7A"},
            {"label": "Nuclear entry", "caption": "Factors enter the nucleus.", "color": "#B7791F"},
            {"label": "Matrix synthesis", "caption": "Matrix is restored.", "color": "#6B46C1"},
        ],
    }

    def _asset(self):
        import base64
        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAFgwJ/lP3GJwAAAABJRU5ErkJggg==")
        return compose.asset_image(png)

    def test_plate_keeps_painted_scene_and_all_stage_labels(self):
        svg = archetypes_render._render_signal_plate(
            self.SPEC, self._asset(), [1280, 900], "rect",
            archetypes_render._palette({}, "signal_pathway"))
        pipeline._check_svg(svg, allow_image=True)
        self.assertEqual(svg.count(b"<image"), 1)
        self.assertGreaterEqual(svg.count(b"<text"), 10)
        for stage in self.SPEC["stages"]:
            self.assertIn(stage["label"].encode(), svg)

    def test_scene_renderer_uses_plate_asset_instead_of_cards(self):
        svg = archetypes_render._render_signal_scene(
            self.SPEC, {"scene": self._asset()}, [1280, 900], "rect",
            archetypes_render._palette({}, "signal_pathway"))
        self.assertEqual(svg.count(b"<image"), 1)
        self.assertIn(b"Ligand release", svg)

    def test_render_kind_reuses_saved_scene_without_compose_name_error(self):
        import io
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (64, 64), "white").save(buf, "PNG")
        png = buf.getvalue()
        with tempfile.TemporaryDirectory() as tmp:
            scene = Path(tmp) / "scene.png"
            scene.write_bytes(png)
            spec = dict(self.SPEC)
            spec["archetype"] = "signal_pathway"
            spec["layout"] = "horizontal_s"
            report = {"archetype": "signal_pathway", "scene_template": str(scene)}
            svg, _, _ = app_server._render_kind(
                "archetype", spec, report, {}, [1280, 900], "rect",
                "Times New Roman", {"illustration_mode": "image"})
        self.assertEqual(svg.count(b"<image"), 1)

    def test_six_stage_plate_keeps_every_label_and_one_painting(self):
        spec = dict(self.SPEC)
        spec["stages"] = self.SPEC["stages"] + [
            {"label": "Feedback to receptor", "caption": "Signal is damped.", "color": "#8B5CF6"}]
        svg = archetypes_render._render_signal_plate(
            spec, self._asset(), [1440, 1080], "rect",
            archetypes_render._palette({}, "signal_pathway"))
        pipeline._check_svg(svg, allow_image=True)
        self.assertEqual(svg.count(b"<image"), 1)
        for stage in spec["stages"]:
            self.assertIn(stage["label"].encode(), svg)
        self.assertGreaterEqual(svg.count(b"<text"), 12)


class LayoutContractTests(unittest.TestCase):
    def test_whole_figure_prompt_uses_author_layout_and_labels(self):
        cfg = {"layout_contract": {
            "title": "EV Signaling", "layout": "vertical",
            "stages": ["Ligand secretion", "Receptor binding", "Matrix synthesis"]}}
        prompt = pipeline._whole_figure_prompt("brief", cfg)
        self.assertIn("Vertical top-to-bottom", prompt)
        self.assertIn("Ligand secretion -> Receptor binding -> Matrix synthesis", prompt)
        self.assertIn("85%", prompt)
        self.assertIn("Do not reuse the same generic circle", prompt)

    def test_reference_reading_order_wins_without_layout_words(self):
        contract = layout_contract.build(
            "配体→受体→胞内级联→核内效应。",
            reference_layout={"layout_analysis": {
                "composition": "one cell on the left, matrix on the right",
                "reading_order": "left-to-right",
                "label_position": "outside left and right edges",
                "visual_hierarchy": "cell scene dominant"}})
        self.assertEqual(contract["archetype"], "signal_pathway")
        self.assertEqual(contract["layout"], "horizontal_s")
        self.assertEqual(contract["figure_style"], "mechanism")
        self.assertIn("outside left", contract["reference_layout"]["label_position"])

    def test_explicit_flowchart_request_keeps_cards(self):
        contract = layout_contract.build(
            "请画成带方框的流程图，三个步骤。",
            reference_layout={"composition": "panels", "reading_order": "left-to-right"})
        self.assertEqual(contract["figure_style"], "flowchart")

    def test_left_heavy_card_is_flagged(self):
        svg = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1600 1000">
          <rect x="100" y="100" width="1200" height="160" rx="16" fill="#fff" stroke="#000"/>
          <text x="130" y="180" font-size="20">Only a short label</text>
        </svg>'''
        findings = review.findings(svg.encode(), None)
        self.assertTrue(any(f["kind"] == "layout" and f["severity"] == "high"
                            for f in findings))


class PipelineLayoutTests(unittest.TestCase):
    SPEC = {
        "title": "Cartilage repair process",
        "subtitle": "From injury to matrix restoration",
        "stages": [
            {"label": "Injury", "caption": "Cartilage is damaged."},
            {"label": "Inflammation", "caption": "Immune cells arrive."},
            {"label": "Repair", "caption": "Chondrocytes rebuild matrix."},
            {"label": "Remodelling", "caption": "Matrix is organised."},
        ],
    }

    def test_plain_process_uses_textbook_strip_not_cards(self):
        svg = archetypes_render._render_pipeline(
            self.SPEC, {}, [1280, 900], "rect",
            archetypes_render._palette({}, "process_pipeline"))
        pipeline._check_svg(svg, allow_image=True)
        # Journal figures carry no in-artwork title; the caption is external.
        self.assertNotIn(b"Cartilage repair process", svg)
        self.assertIn(b"Injury", svg)
        for stage in self.SPEC["stages"]:
            self.assertIn(stage["label"].encode(), svg)
        self.assertNotIn(b'rx="14"', svg)

    def test_explicit_flowchart_keeps_card_grid(self):
        spec = dict(self.SPEC)
        spec["_flowchart"] = True
        svg = archetypes_render._render_pipeline(
            spec, {}, [1280, 900], "rect",
            archetypes_render._palette({}, "process_pipeline"))
        self.assertIn(b'rx="14"', svg)


class ComposeLayoutTests(unittest.TestCase):
    @staticmethod
    def _scene(count, labels=None):
        labels = labels or [f"Entity {i + 1}" for i in range(count)]
        return {
            "canvas": [1280, 900], "title": "Tumor microenvironment",
            "palette": {"text": "#102a33", "primary": "#3F5A6B"},
            "entities": [{"id": f"e{i}", "label": labels[i]} for i in range(count)],
            "connectors": [], "labels": [],
        }

    def test_three_entities_use_one_centered_row(self):
        scene = self._scene(3)
        compose.arrange(scene)
        boxes = [e["bbox"] for e in scene["entities"]]
        self.assertEqual(len({round(box[1]) for box in boxes}), 1)
        self.assertLess(boxes[0][0], boxes[1][0])
        left = boxes[0][0]
        right = boxes[-1][0] + boxes[-1][2]
        self.assertAlmostEqual((left + right) / 2, 640, delta=1)

    def test_partial_second_row_is_centered(self):
        scene = self._scene(5)
        compose.arrange(scene)
        second = [e["bbox"] for e in scene["entities"][3:]]
        self.assertEqual(len(second), 2)
        left = second[0][0]
        right = second[-1][0] + second[-1][2]
        self.assertAlmostEqual((left + right) / 2, 640, delta=1)

    def test_caption_band_does_not_touch_next_row(self):
        scene = self._scene(5)
        compose.arrange(scene)
        top = scene["entities"][0]["bbox"]
        bottom = scene["entities"][3]["bbox"]
        self.assertGreaterEqual(bottom[1] - (top[1] + top[3]), 44)

    def test_labels_stay_inside_canvas_and_do_not_overlap(self):
        labels = ["Cancer-associated fibroblast (CAF)", "Tumor cell",
                  "M2 macrophage", "Collagen-rich ECM", "Endothelial cell"]
        scene = self._scene(5, labels)
        compose.arrange(scene)
        asset = ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100">'
                 '<circle cx="50" cy="50" r="40" fill="#999"/></svg>')
        assets = {e["id"]: (asset, 100.0, 100.0) for e in scene["entities"]}
        svg = compose.compose_svg(scene, assets, "Times New Roman")
        import xml.etree.ElementTree as ET
        boxes = review._text_boxes(ET.fromstring(svg))
        self.assertGreaterEqual(len(boxes), 5)
        for box in boxes:
            self.assertGreaterEqual(box["x"], 0)
            self.assertGreaterEqual(box["y"], 0)
            self.assertLessEqual(box["x"] + box["w"], 1280)
            self.assertLessEqual(box["y"] + box["h"], 900)
        for i, first in enumerate(boxes):
            for second in boxes[i + 1:]:
                width = min(first["x"] + first["w"], second["x"] + second["w"]) - max(first["x"], second["x"])
                height = min(first["y"] + first["h"], second["y"] + second["h"]) - max(first["y"], second["y"])
                self.assertFalse(width > 1 and height > 1,
                                 f"{first['text']} overlaps {second['text']}")


class ReviewCrashTests(unittest.TestCase):
    def test_framework_findings_returns_a_list(self):
        svg = '''<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 1000 700">
          <defs><clipPath id="c"><rect x="100" y="130" width="800" height="100"/></clipPath></defs>
          <text x="500" y="55" text-anchor="middle" font-size="30">Framework Title</text>
          <g clip-path="url(#c)"><rect x="100" y="130" width="800" height="100" fill="#eee"/></g>
        </svg>'''
        result = review.findings(svg.encode(), None)
        self.assertIsInstance(result, list)

    def test_framework_escapes_ampersands_in_wrapped_text(self):
        spec = {
            "title": "A & B", "subtitle": "C & D",
            "priors_items": [{"title": "X & Y", "desc": "P & Q"}],
            "latent_items": [{"title": "L & M", "desc": "N & O"}],
            "ai_items": [{"title": "G & H", "desc": "I & J"}],
            "sensing_modalities": [], "embodied_entities": [],
        }
        svg = framework.build_framework_svg(spec, {}, "Times New Roman",
                                            canvas=[1200, 800])
        pipeline._check_svg(svg, allow_image=True)


class VisionGuardTests(unittest.TestCase):
    def test_spatial_composition_high_is_downgraded(self):
        svg = (b'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 800 600">'
               b'<text x="40" y="60" font-size="24">Tumor cell</text>'
               b'<text x="300" y="60" font-size="24">CAF</text>'
               b'<text x="560" y="60" font-size="24">ECM</text>'
               b'<text x="40" y="500" font-size="24">Macrophage</text></svg>')
        findings = [{"kind": "vision", "severity": "high",
                     "message": "空间构成与参考描述不符"}]
        kept = model_review._geometry_guard(findings, svg)
        self.assertEqual(kept[0]["severity"], "medium")


class ExportAuditTests(unittest.TestCase):
    def test_svg_audit_counts_images_and_texts(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "figure.svg"
            path.write_bytes(b'<svg xmlns="http://www.w3.org/2000/svg">'
                             b'<image href="data:image/png;base64,AA"/>'
                             b'<image href="data:image/png;base64,BB"/>'
                             b'<text>A</text><text>B</text></svg>')
            audit = app_server._svg_audit(path)
            self.assertEqual(audit["images"], 2)
            self.assertEqual(audit["texts"], 2)

    def test_pptx_audit_counts_pictures_texts_and_media(self):
        import zipfile
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "figure.pptx"
            with zipfile.ZipFile(path, "w") as package:
                package.writestr("ppt/slides/slide1.xml",
                                 "<p:pic/><p:pic/><a:t>x</a:t><a:t>y</a:t><a:t>z</a:t>")
                package.writestr("ppt/media/image1.png", b"png")
            audit = app_server._pptx_audit(path)
            self.assertEqual(audit["pictures"], 2)
            self.assertEqual(audit["texts"], 3)
            self.assertEqual(audit["media"], 1)


class VisionModelSelectionTests(unittest.TestCase):
    def _config(self, model="local-vl"):
        return {"vision": {"base_url": "http://vision.test/v1", "model": model,
                           "api_key": ""}}

    def test_config_model_is_used_first(self):
        from unittest import mock
        calls = []
        def fake_post(url, key, payload, timeout=0):
            calls.append(payload["model"])
            return {"choices": [{"message": {"content":
                    '{"verdict":"pass","score":100,"issues":[]}'}}]}
        with mock.patch.object(model_review, "_image_data_url", return_value="data:image/png;base64,AA"), \
             mock.patch.object(pipeline, "_post_json", side_effect=fake_post), \
             mock.patch.object(pipeline, "_endpoint", return_value="http://vision.test/v1/chat/completions"), \
             mock.patch.object(pipeline, "_api_key", return_value=""):
            findings, meta = model_review.review_figure(b"<svg/>", self._config())
        self.assertEqual(calls, ["local-vl"])
        self.assertEqual(findings, [])
        self.assertFalse(meta["fallback"])
        self.assertEqual(meta["configured_model"], "local-vl")

    def test_config_failure_falls_back_to_critic(self):
        from unittest import mock
        calls = []
        def fake_post(url, key, payload, timeout=0):
            calls.append(payload["model"])
            if payload["model"] == "local-vl":
                raise RuntimeError("boom")
            return {"choices": [{"message": {"content":
                    '{"verdict":"pass","score":90,"issues":[]}'}}]}
        with mock.patch.object(model_review, "_image_data_url", return_value="data:image/png;base64,AA"), \
             mock.patch.object(pipeline, "_post_json", side_effect=fake_post), \
             mock.patch.object(pipeline, "_endpoint", return_value="http://vision.test/v1/chat/completions"), \
             mock.patch.object(pipeline, "_api_key", return_value=""):
            _, meta = model_review.review_figure(b"<svg/>", self._config())
        self.assertEqual(calls[0], "local-vl")
        self.assertEqual(calls[-1], model_review.VISION_MODEL)
        self.assertTrue(meta["fallback"])
        self.assertEqual(meta["model"], model_review.VISION_MODEL)


class PptRegressionTests(unittest.TestCase):
    def build(self, svg: str):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        svg_path = Path(tmp.name) / "figure.svg"
        svg_path.write_text(svg, encoding="utf-8")
        return ppt_native.build(svg_path, Path(tmp.name) / "out", font="Times New Roman")

    def test_textpath_labels_are_written(self):
        svg = '''<svg xmlns="http://www.w3.org/2000/svg" width="400" height="200"
            viewBox="0 0 400 200">
          <defs><path id="arc" d="M 50 120 A 150 150 0 0 1 350 120"/></defs>
          <text font-size="20" fill="#102a33" data-ppt-x="200" data-ppt-y="80"
                data-ppt-rot="0" data-ppt-anchor="middle"><textPath href="#arc"
                startOffset="50%" text-anchor="middle">Arc Label</textPath></text>
        </svg>'''
        report = self.build(svg)
        self.assertEqual(report["texts"], 1)
        from pptx import Presentation
        prs = Presentation(report["pptx"])
        texts = [shape.text_frame.text for slide in prs.slides for shape in slide.shapes
                 if shape.has_text_frame]
        self.assertIn("Arc Label", texts)

    def test_scaled_group_scales_font_size(self):
        svg = '''<svg xmlns="http://www.w3.org/2000/svg" width="400" height="200"
            viewBox="0 0 400 200">
          <g transform="scale(0.5)"><text x="100" y="100" font-size="40"
            fill="#102a33">Scaled</text></g>
        </svg>'''
        report = self.build(svg)
        from pptx import Presentation
        from pptx.util import Pt
        prs = Presentation(report["pptx"])
        runs = [run for slide in prs.slides for shape in slide.shapes
                if shape.has_text_frame for para in shape.text_frame.paragraphs
                for run in para.runs]
        self.assertTrue(runs)
        self.assertAlmostEqual(float(runs[0].font.size.pt), 15.0, delta=0.6)

    def test_multiline_text_compares_clean(self):
        svg = '''<svg xmlns="http://www.w3.org/2000/svg" width="400" height="200"
            viewBox="0 0 400 200">
          <text x="40" y="80" font-size="18" fill="#102a33">
            <tspan x="40" y="80">Latent Inference &amp;</tspan>
            <tspan x="40" y="110">Compliant Actuation</tspan>
          </text>
        </svg>'''
        report = self.build(svg)
        findings = review.compare_texts(svg.encode(), report["pptx"])
        self.assertEqual(findings, [])


class ImageOutageTests(unittest.TestCase):
    """An unreachable image model must fail loudly, never yield a blank page."""

    CONFIG = {
        "text": {"base_url": "http://127.0.0.1:9/v1", "model": "t", "api_key": "k"},
        "image": {"base_url": "http://127.0.0.1:9/v1", "model": "i", "api_key": "k"},
        "render_mode": "compose",
    }

    def test_http_5xx_becomes_image_service_error(self):
        from unittest import mock
        with mock.patch.object(pipeline, "_post_json",
                               side_effect=RuntimeError("HTTP 503: upstream down")):
            with self.assertRaises(pipeline.ImageServiceError):
                pipeline.generate_raw("a cell", self.CONFIG)

    def test_compose_outage_is_not_masked_by_plan_fallback(self):
        from unittest import mock
        with mock.patch.object(compose, "compose_render",
                               side_effect=pipeline.ImageServiceError("生图服务不可用：503")):
            with mock.patch.object(pipeline, "plan_to_svg") as plan:
                with self.assertRaises(pipeline.ImageServiceError):
                    pipeline.auto_render("a cell", self.CONFIG)
                plan.assert_not_called()


class AutoFixTraceTests(unittest.TestCase):
    def test_spec_diff_names_the_changed_stage(self):
        before = {"stages": [{"label": "A", "prompt": "old"}]}
        after = {"stages": [{"label": "A", "prompt": "new"}]}
        actions = app_server._spec_change_actions(before, after)
        self.assertTrue(actions)
        self.assertIn("stage 1 prompt", actions[0])
        self.assertIn("old", actions[0])
        self.assertIn("new", actions[0])

    def test_plate_heuristic_makes_a_visible_spec_change(self):
        spec = {"stages": [{"label": "A", "caption": "a"}]}
        actions = model_review.heuristic_fix(
            spec, [{"severity": "high", "category": "legibility",
                    "message": "阶段标签字号偏小，看不清"}])
        self.assertTrue(actions)
        self.assertGreater(spec.get("_plate_label_scale", 1.0), 1.0)


class ReferenceQuestionTests(unittest.TestCase):
    def test_layout_and_content_questions_come_first(self):
        dna = {
            "reference_summary": "单幅连续场景，主体居中，标签在画外。",
            "layout_analysis": {
                "composition": "continuous_scene",
                "reading_order": "left-to-right",
                "label_position": "outside",
            },
        }
        cards = assist_guide._attach_reference(
            {"summary": "下面 6 个问题把这个排版固定下来。", "questions": [
                {"id": "archetype", "title": "图型", "options": [
                    {"value": "a", "label": "A"}, {"value": "b", "label": "B"}]},
            ]}, dna)
        ids = [q["id"] for q in cards["questions"]]
        self.assertEqual(ids[:2], ["reference_layout", "reference_content"])
        self.assertIn("连续场景", cards["questions"][0]["title"])
        self.assertIn("请确认下面 3 项", cards["summary"])


if __name__ == "__main__":
    unittest.main()
