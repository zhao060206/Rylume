import io
import tempfile
import zipfile
from pathlib import Path

import fitz
import docx
import pptx
import openpyxl
import pytest

from document_conversion import (
    convert_document,
    check_container,
    ConversionError,
    DocumentCancelled,
    office_path,
    MAX_UPLOAD,
    SUPPORTED,
)


def test_pdf_conversion_with_locators(tmp_path):
    pdf_path = tmp_path / "test.pdf"
    doc = fitz.open()
    p1 = doc.new_page()
    p1.insert_text((50, 50), "First page: EGFR signaling pathway activates downstream MAPK cascade.")
    p2 = doc.new_page()
    p2.insert_text((50, 50), "Second page: STAT3 translocates to the nucleus to induce transcription.")
    doc.save(str(pdf_path))
    doc.close()

    result = convert_document(pdf_path)
    assert result["pages"] == 2
    assert len(result["sections"]) == 2
    assert "EGFR" in result["markdown"]
    assert "STAT3" in result["markdown"]
    assert "第 1 页" in result["sections"][0]["locator"]
    assert "第 2 页" in result["sections"][1]["locator"]
    assert "PyMuPDF 页码定位" in result["converter"]


def test_docx_conversion_with_headings(tmp_path):
    docx_path = tmp_path / "test.docx"
    doc = docx.Document()
    doc.add_heading("Experimental Design", level=1)
    doc.add_paragraph("Sample incubation with primary antibody for 2 hours at room temperature.")
    doc.add_heading("Detection", level=1)
    doc.add_paragraph("Fluorophore conjugated secondary antibody incubated in the dark.")
    doc.save(str(docx_path))

    result = convert_document(docx_path)
    assert len(result["sections"]) >= 2
    assert "Experimental Design" in result["markdown"]
    assert "Detection" in result["markdown"]
    assert "Microsoft MarkItDown" in result["converter"]


def test_pptx_conversion_with_slide_numbers(tmp_path):
    pptx_path = tmp_path / "test.pptx"
    prs = pptx.Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Mechanism Overview"
    s1.placeholders[1].text = "Receptor dimerization initiates signaling."
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Target Genes"
    s2.placeholders[1].text = "Cyclin D1 and c-Myc expression upregulated."
    prs.save(str(pptx_path))

    result = convert_document(pptx_path)
    assert result["pages"] >= 2
    assert "Mechanism Overview" in result["markdown"]
    assert "Target Genes" in result["markdown"]
    assert any("幻灯片" in s["locator"] for s in result["sections"])


def test_xlsx_conversion(tmp_path):
    xlsx_path = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Gene", "FoldChange", "PValue"])
    ws.append(["STAT3", "2.4", "0.001"])
    ws.append(["MYC", "3.1", "0.0004"])
    wb.save(str(xlsx_path))

    result = convert_document(xlsx_path)
    assert "STAT3" in result["markdown"]
    assert "PValue" in result["markdown"]
    assert "工作表 Results" in result["markdown"]
    assert "openpyxl" in result["converter"]


def test_plain_text_and_markdown_conversion(tmp_path):
    txt_path = tmp_path / "notes.txt"
    txt_path.write_text("Cell lysis buffer contains protease inhibitors and detergent.", encoding="utf-8")
    res_txt = convert_document(txt_path)
    assert "Cell lysis buffer" in res_txt["markdown"]

    md_path = tmp_path / "protocol.md"
    md_path.write_text("# Workflow\n\n1. Centrifuge at 12000 rpm\n2. Collect supernatant", encoding="utf-8")
    res_md = convert_document(md_path)
    assert "Centrifuge at 12000 rpm" in res_md["markdown"]


def test_csv_json_html_conversion(tmp_path):
    csv_path = tmp_path / "data.csv"
    csv_path.write_text("Time,OD600\n0,0.1\n1,0.25\n2,0.6", encoding="utf-8")
    res_csv = convert_document(csv_path)
    assert "OD600" in res_csv["markdown"]

    json_path = tmp_path / "config.json"
    json_path.write_text('{"experiment": "western", "target": "p-AKT"}', encoding="utf-8")
    res_json = convert_document(json_path)
    assert "p-AKT" in res_json["markdown"]

    html_path = tmp_path / "page.html"
    html_path.write_text("<html><body><h1>Summary</h1><p>Membrane trafficking mechanism.</p></body></html>", encoding="utf-8")
    res_html = convert_document(html_path)
    assert "Membrane trafficking" in res_html["markdown"]


def test_unsupported_and_corrupt_files(tmp_path):
    exe_path = tmp_path / "program.exe"
    exe_path.write_bytes(b"MZfake")
    with pytest.raises(ConversionError, match="格式"):
        convert_document(exe_path)

    bad_pdf = tmp_path / "broken.pdf"
    bad_pdf.write_bytes(b"NOT A REAL PDF")
    with pytest.raises(ConversionError, match="PDF"):
        convert_document(bad_pdf)

    bad_docx = tmp_path / "broken.docx"
    bad_docx.write_bytes(b"PK\x03\x04corrupted zip")
    with pytest.raises(ConversionError, match="损坏"):
        convert_document(bad_docx)

    # Valid zip but missing word/document.xml
    mismatch_docx = tmp_path / "fake.docx"
    with zipfile.ZipFile(mismatch_docx, "w") as z:
        z.writestr("test.txt", "hello")
    with pytest.raises(ConversionError, match="扩展名不一致"):
        convert_document(mismatch_docx)


def test_empty_file_raises_conversion_error(tmp_path):
    empty_txt = tmp_path / "empty.txt"
    empty_txt.write_text("", encoding="utf-8")
    with pytest.raises(ConversionError, match="可读文字"):
        convert_document(empty_txt)


def test_cancelled_callback_aborts_conversion(tmp_path):
    txt_path = tmp_path / "sample.txt"
    txt_path.write_text("Valid text content for testing cancellation behavior.", encoding="utf-8")
    with pytest.raises(DocumentCancelled):
        convert_document(txt_path, cancelled=lambda: True)


def test_long_section_splitting(tmp_path):
    long_txt = tmp_path / "long.txt"
    # Create 25,000 characters of text without headings
    long_content = "Extracellular ligand binding activates receptor kinase domain. " * 400
    long_txt.write_text(long_content, encoding="utf-8")
    res = convert_document(long_txt)
    # Ensure sections are split into chunks <= 10000 characters
    assert len(res["sections"]) >= 3
    for s in res["sections"]:
        assert len(s["text"]) <= 10000
    assert sum(len(s["text"]) for s in res["sections"]) == len(long_content.strip())


def test_encrypted_pdf_rejected(tmp_path):
    path=tmp_path/'encrypted.pdf'
    doc=fitz.open();doc.new_page().insert_text((50,70),'Synthetic protocol text')
    doc.save(path,encryption=fitz.PDF_ENCRYPT_AES_256,user_pw='fixture-user',owner_pw='fixture-owner');doc.close()
    with pytest.raises(ConversionError,match='加密'):
        convert_document(path)
