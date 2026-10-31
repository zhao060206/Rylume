"""Local-only Markdown conversion with explicit page/slide provenance."""
from __future__ import annotations

import io
import os
import re
import shutil
import subprocess
import tempfile
import time
import zipfile
from pathlib import Path

MAX_UPLOAD = 16 * 1024 * 1024
MAX_MARKDOWN = 800_000
SUPPORTED = {'.pdf', '.pptx', '.docx', '.txt', '.md', '.markdown', '.html', '.htm', '.csv', '.xlsx', '.json', '.xml', '.rtf', '.doc', '.ppt', '.odt', '.odp'}
LEGACY = {'.doc': 'docx', '.ppt': 'pptx', '.rtf': 'docx', '.odt': 'docx', '.odp': 'pptx'}


class ConversionError(ValueError):
    pass


class DocumentCancelled(Exception):
    pass


def office_path():
    candidate = shutil.which('soffice')
    if candidate:
        return Path(candidate)
    for variable in ('PROGRAMFILES', 'PROGRAMFILES(X86)'):
        root = os.environ.get(variable)
        if root and (p := Path(root) / 'LibreOffice/program/soffice.exe').is_file():
            return p
    return None


def check_container(path):
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ConversionError('暂不支持此格式；请使用 PDF、PPTX、DOCX、文本或表格文件。')
    if path.stat().st_size > MAX_UPLOAD:
        raise ConversionError('单个文件请小于 16 MB。')
    head = path.read_bytes()[:8]
    if ext == '.pdf' and not head.startswith(b'%PDF-'):
        raise ConversionError('文件内容不是有效 PDF。')
    if ext in ('.docx', '.pptx', '.xlsx', '.odt', '.odp'):
        try:
            with zipfile.ZipFile(path) as archive:
                infos = archive.infolist()
                if len(infos) > 12000 or sum(x.file_size for x in infos) > 100 * 1024 * 1024:
                    raise ConversionError('文档解压后过大，请移除大型嵌入附件后再上传。')
                if any(x.file_size > 32 * 1024 * 1024 or x.file_size / max(x.compress_size, 1) > 1500 for x in infos):
                    raise ConversionError('文档包含过大的压缩内容。')
                required = {'.docx': 'word/document.xml', '.pptx': 'ppt/presentation.xml', '.xlsx': 'xl/workbook.xml', '.odt': 'content.xml', '.odp': 'content.xml'}[ext]
                if required not in archive.namelist():
                    raise ConversionError('文件内容与扩展名不一致。')
        except zipfile.BadZipFile as error:
            raise ConversionError('Office 文件已损坏或扩展名不正确。') from error
    if ext in ('.doc', '.ppt') and head != b'\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1':
        raise ConversionError('旧版 Office 文件内容不正确；请另存为 DOCX 或 PPTX。')


def _clean(text):
    text = str(text).replace('\x00', '')
    text = re.sub(r'!\[([^\]]*)\]\(data:[^)]+\)', r'[图片：\1]', text)
    return text.strip()


def _ocr_page(page):
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR
    pix = page.get_pixmap(matrix=__import__('fitz').Matrix(1.4, 1.4), alpha=False)
    image = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)
    result, _ = RapidOCR()(image[:, :, :3])
    return '\n'.join(row[1] for row in result or [])


def _office_convert(path, folder):
    executable = office_path()
    if not executable:
        raise ConversionError('旧版 PPT、Word 与 RTF 需要 LibreOffice。请安装 LibreOffice，或在原软件另存为 PPTX、DOCX 后上传。')
    profile = (folder / 'office-profile').resolve().as_uri()
    command = [str(executable), '-env:UserInstallation=' + profile, '--headless', '--norestore', '--convert-to', LEGACY[path.suffix.lower()], '--outdir', str(folder), str(path)]
    try:
        subprocess.run(command, capture_output=True, timeout=90, check=True, creationflags=getattr(subprocess, 'CREATE_NO_WINDOW', 0))
    except (subprocess.SubprocessError, OSError) as error:
        raise ConversionError('LibreOffice 转换未完成，请将文件另存为 PPTX 或 DOCX 后重试。') from error
    output = folder / (path.stem + '.' + LEGACY[path.suffix.lower()])
    if not output.is_file():
        raise ConversionError('LibreOffice 没有生成转换文件。')
    return output


def convert_document(path, *, cancelled=lambda: False, progress=lambda *_: None, ocr=False):
    from markitdown import MarkItDown, StreamInfo
    path = Path(path)
    check_container(path)
    started = time.monotonic()
    warnings, parts, scanned = [], [], []
    engine = MarkItDown(enable_plugins=False)

    def check():
        if cancelled():
            raise DocumentCancelled()
        if time.monotonic() - started > 150:
            raise ConversionError('转换超过 150 秒，请拆分文件后重试。')

    def add(locator, text):
        check()
        text = _clean(text)
        if text:
            parts.append({'locator': locator, 'text': text})
        if sum(len(p['text']) for p in parts) > MAX_MARKDOWN:
            raise ConversionError('转换文字超过 80 万字，请拆分文档；未截断或发送不完整内容。')

    with tempfile.TemporaryDirectory(prefix='md-convert-', dir=path.parent) as temporary:
        folder = Path(temporary)
        source = _office_convert(path, folder) if path.suffix.lower() in LEGACY else path
        if source != path:
            check_container(source)
        check()
        ext = source.suffix.lower()
        converter = 'Microsoft MarkItDown 0.1.9' + (' + LibreOffice' if source != path else '')
        if ext == '.pdf':
            import fitz
            with fitz.open(source) as pdf:
                if pdf.needs_pass:
                    raise ConversionError('PDF 已加密，请解密后上传。')
                if len(pdf) > 200:
                    raise ConversionError('单个 PDF 最多 200 页，请按章节拆分。')
                for index, page in enumerate(pdf):
                    check()
                    progress(round(index * 90 / max(len(pdf), 1)), f'正在转换第 {index + 1} / {len(pdf)} 页')
                    raw = page.get_text(sort=True).strip()
                    if len(raw) < 15 and page.get_images():
                        scanned.append(index + 1)
                        if ocr and len(scanned) <= 20:
                            raw = _ocr_page(page)
                        else:
                            raw = ''
                    if raw:
                        single = fitz.open()
                        single.insert_pdf(pdf, from_page=index, to_page=index)
                        data = single.tobytes()
                        single.close()
                        try:
                            md = engine.convert_stream(io.BytesIO(data), stream_info=StreamInfo(extension='.pdf')).text_content
                        except Exception:
                            md = raw
                            if '部分页使用 PyMuPDF 提取' not in warnings:
                                warnings.append('部分页使用 PyMuPDF 提取')
                        add(f'第 {index + 1} 页', md if len(md.strip()) >= 15 else raw)
                    else:
                        note = '，可启用本机 OCR 重试' if not ocr else ('，本次 OCR 最多识别 20 个扫描页，请拆分后重试' if len(scanned) > 20 else '，请核对扫描质量')
                        warnings.append(f'第 {index + 1} 页未提取到文字' + note)
                page_count = len(pdf)
            converter += ' + PyMuPDF 页码定位' + (' + RapidOCR' if ocr and scanned else '')
        elif ext == '.xlsx':
            from openpyxl import load_workbook
            workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
            try:
                for sheet in workbook:
                    rows = []
                    for index, row in enumerate(sheet.iter_rows(values_only=True)):
                        check()
                        if index >= 10000:
                            raise ConversionError('表格超过 10000 行，请选择相关数据后上传。')
                        cells = [str(v if v is not None else '').replace('|', '\\|').replace('\n', ' ') for v in row[:80]]
                        if len(row) > 80:
                            raise ConversionError('表格超过 80 列，请拆分后上传。')
                        if any(cells):
                            rows.append('| ' + ' | '.join(cells) + ' |')
                    if rows:
                        rows.insert(1, '| ' + ' | '.join('---' for _ in rows[0].split('|')[1:-1]) + ' |')
                        add('工作表 ' + sheet.title, '\n'.join(rows))
            finally:
                workbook.close()
            converter = 'openpyxl → Markdown'
            page_count = len(parts)
        else:
            result = _clean(engine.convert_local(source).text_content)
            if ext == '.pptx':
                slides = re.split(r'(?=<!--\s*Slide number:\s*\d+\s*-->)', result)
                for index, slide in enumerate(s for s in slides if s.strip()):
                    match = re.search(r'Slide number:\s*(\d+)', slide)
                    add('第 ' + (match.group(1) if match else str(index + 1)) + ' 张幻灯片', slide)
            else:
                blocks = re.split(r'\n(?=#{1,4}\s)', result)
                for index, block in enumerate(blocks):
                    heading = block.split('\n', 1)[0].lstrip('# ').strip()
                    add((heading[:80] if block.startswith('#') else f'正文第 {index + 1} 节'), block)
            page_count = len(parts)
        check()
    if not parts:
        raise ConversionError('没有提取到可读文字。扫描 PDF 可启用本机 OCR 重试；仅有图片的幻灯片请补充文字说明。')
    sections = []
    for index, part in enumerate(parts):
        # Long sections are split without silently dropping any text.
        text = part['text']
        for start in range(0, len(text), 10000):
            sections.append({'key': f's{len(sections) + 1:04}', 'locator': part['locator'], 'text': text[start:start + 10000]})
    markdown = '\n\n'.join('## ' + p['locator'] + '\n\n' + p['text'] for p in parts)
    return dict(markdown=markdown, sections=sections, converter=converter, warnings=warnings, pages=page_count, scanned_pages=scanned)
